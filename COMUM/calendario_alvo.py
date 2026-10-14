"""Busca o ALVO do dia (VALOR_EB) na planilha de calendario de planejamento
de baixas -- X:\\...\\MKP\\CALENDARIO\\CALENDARIO_<ano>\\Calendario_Planejamento
Baixas_<ano>.xlsx, aba "Visao_<MES>_<ano>". O calendario ja tem o valor
esperado por LOJA/DATA lancado com antecedencia; o nome da LOJA vem de
CANAIS[canal]["loja_calendario"] em canais.py (fonte unica dessa config
por canal).

Usado pra validar que RB + REBATE + AJUSTE (fila VERIFICAR) fecha com o
alvo do dia -- pedido do Everton, ver dashboard.py.
"""

import glob
import os
import sys
import unicodedata

import openpyxl

sys.path.insert(0, os.path.join(os.path.dirname(__file__), "..", "CONCILIACAO"))
from canais import CANAIS

PASTA_CALENDARIO = r"X:\109_GC_CTR_FINANCEIRO CONTAS A RECEBER\MKP\CALENDÁRIO"

MESES = [
    "JANEIRO", "FEVEREIRO", "MARCO", "ABRIL", "MAIO", "JUNHO",
    "JULHO", "AGOSTO", "SETEMBRO", "OUTUBRO", "NOVEMBRO", "DEZEMBRO",
]


def _sem_acento(texto):
    nfkd = unicodedata.normalize("NFKD", texto)
    return "".join(c for c in nfkd if not unicodedata.combining(c))


def _normalizar(texto):
    return _sem_acento(texto).upper().replace(" ", "").replace("\xa0", "")


def _achar_planilha_calendario(ano):
    padrao = os.path.join(PASTA_CALENDARIO, f"CALENDÁRIO_{ano}", "Calendário_Planejamento Baixas*.xlsx")
    achados = sorted(f for f in glob.glob(padrao) if not os.path.basename(f).startswith("~$"))
    return achados[-1] if achados else None


def _achar_aba_do_mes(wb, mes, ano):
    alvo = _normalizar(f"VISAO_{mes}_{ano}")
    for nome in wb.sheetnames:
        if _normalizar(nome) == alvo:
            return wb[nome]
    return None


def buscar_alvo(canal, data_ref):
    """data_ref: 'dd/mm/aaaa' (mesmo formato de parser_rb.metadata_rb).
    Retorna o VALOR_EB do calendario pro canal/dia, ou None se nao achar
    planilha, aba ou linha (mes ainda nao cadastrado, canal fora do
    calendario etc -- nao e erro, so significa "sem alvo pra comparar")."""
    return buscar_alvo_loja(CANAIS.get(canal, {}).get("loja_calendario"), data_ref)


def buscar_alvo_loja(loja, data_ref):
    """Mesma busca de buscar_alvo, mas direto pelo nome da LOJA no calendario
    (usado pelo koncili.py, que confere canais que nao estao no CANAIS)."""
    if not loja or not data_ref:
        return None

    try:
        dia, mes_num, ano = (int(p) for p in data_ref.split("/"))
    except ValueError:
        return None

    caminho = _achar_planilha_calendario(ano)
    if not caminho:
        return None

    wb = openpyxl.load_workbook(caminho, data_only=True)
    ws = _achar_aba_do_mes(wb, MESES[mes_num - 1], ano)
    if ws is None:
        return None

    header = [c.value for c in ws[1]]
    obrigatorias = ("LOJA", "DATA", "VALOR_EB")
    if not all(col in header for col in obrigatorias):
        return None
    idx = {h: i for i, h in enumerate(header)}
    idx_marketplace = idx.get("MARKETPLACE")

    for row in ws.iter_rows(min_row=2, values_only=True):
        if idx_marketplace is not None and row[idx_marketplace] != "MERCADO":
            continue
        loja_linha = _normalizar(str(row[idx["LOJA"]] or ""))
        if loja_linha != _normalizar(loja):
            continue
        data_linha = row[idx["DATA"]]
        if data_linha is None:
            continue
        if (data_linha.day, data_linha.month, data_linha.year) == (dia, mes_num, ano):
            valor = row[idx["VALOR_EB"]]
            return float(valor) if valor is not None else None
    return None


if __name__ == "__main__":
    alvo = buscar_alvo("FULL", "01/09/2026")
    assert alvo is not None, "nao achou o alvo FULL 01/09/2026 no calendario"
    assert abs(alvo - 102312.33) < 0.01, f"alvo inesperado: {alvo}"
    assert buscar_alvo("FULL", "31/02/2026") is None  # dia que nao existe -> None, sem quebrar
    assert buscar_alvo("XPTO", "01/09/2026") is None  # canal fora do calendario -> None
    print("calendario_alvo.py: ok")
