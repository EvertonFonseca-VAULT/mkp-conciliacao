"""Empilhado BARRA CRED (gohorse -- regra provisoria, estruturar depois, ver conversa 2026-10-13).

Fonte: X:\\...\\MKP - BARRA CRED\\<CANAL>\\<ano>\\<mes>\\BARRA CRED <CANAL> DD.MM.AAAA_AB.xlsx
  aba Omnik (CENTER) / Planilha1 (WEB) -- layout fixo de 16 colunas (Pedido..Obs), ultima coluna
  (Obs) descartada. So entram linhas com "Status do Pedido" == "Entregue".

Destino: CENTER_BARRA_CRED.xlsx / WEB_BARRA_CRED.xlsx, aba de dados -- layout fixo de 18 colunas.
  IMPORTANTE: CENTER e WEB NAO seguem a mesma convencao de tipo por coluna (arquivos mantidos
  a mao, convencao foi mudando com o tempo) -- conferido direto contra as ultimas linhas
  autenticas de cada um (nao as mais antigas, que tem outra convencao ainda):
    - Chave_Primaria: formula =CONCATENATE(Q<linha>,"_","<SUFIXO>") nos dois (NAO literal).
    - CENTER: "Data de Criacao do Lancamento" e "Data do Status do Pedido" ficam como TEXTO
      (string "DD/MM/AAAA" crua, igual a fonte) -- "Pedido Site"/"Status Lancamento"/
      "Data do Status Lancamento"/"Descricao" ficam preenchidos normalmente.
    - WEB: "Data de Criacao do Lancamento" vira DATETIME; "Data do Status do Pedido" vira
      SERIAL do Excel (inteiro, nao datetime nem texto); "Pedido Site", "Status Lancamento",
      "Data do Status Lancamento" e "Descricao" ficam sempre em branco (convencao atual do
      arquivo derruba esses campos -- a fonte tem o dado, o destino so nao usa mais).
  Numero de formato (number_format) de cada coluna: copiado da ultima linha autentica.

Dedupe: chave natural (Pedido, Tipo, Parcela, Data do Status Lancamento, Valor) no CENTER;
  no WEB essa "Data do Status Lancamento" e sempre None no destino (convencao do arquivo),
  entao a chave fica so (Pedido, Tipo, Parcela, Valor) -- testada, 1 colisao em 2025 linhas
  (pre-existente, nao criada por este script; aceito, gohorse).

Uso: python empilhar_barra_cred.py CENTER
     python empilhar_barra_cred.py WEB
"""
import datetime
import os
import shutil
import sys
from copy import copy

import openpyxl
from openpyxl.utils.datetime import to_excel

PASTA_FONTE_BASE = r"X:\109_GC_CTR_FINANCEIRO CONTAS A RECEBER\MKP\MKP - BARRA CRED"
PASTA_DESTINO = (r"X:\109_GC_CTR_FINANCEIRO CONTAS A RECEBER\MKP\ARQUIVOS EMPILHADOS DE BAIXAS"
                 r"\ARQUIVOS REPASSES COM CHAVES PRIMÁRIAS")
PASTA_BACKUP = r"C:\BACKUP EMPILHADO"

# indices (0-based) do layout fixo de 16 colunas da fonte (Omnik/Planilha1)
F_PEDIDO, F_MKT, F_SELLER, F_DT_CRIACAO, F_PEDIDO_SITE, F_STATUS_PEDIDO, F_DT_STATUS_PEDIDO, \
    F_TIPO, F_STATUS_LANC, F_DT_STATUS_LANC, F_DESCRICAO, F_VALOR, F_DT_PREVISTA, F_ID_INTEGRADOR, \
    F_PARCELA, F_OBS = range(16)

# indices (0-based) do layout fixo de 18 colunas do destino
D_DATA_REPASSE, D_MKT, D_SELLER, D_DT_CRIACAO, D_PEDIDO, D_PEDIDO_SITE, D_STATUS_PEDIDO, \
    D_DT_STATUS_PEDIDO, D_TIPO, D_STATUS_LANC, D_DT_STATUS_LANC, D_DESCRICAO, D_VALOR, \
    D_DT_PREVISTA, D_ID_INTEGRADOR, D_PARCELA, D_SEQUENCIA, D_CHAVE = range(18)

CANAIS = {
    "CENTER": dict(aba_fonte="Omnik", arquivo="CENTER_BARRA_CRED.xlsx", aba_destino="CENTER",
                   sufixo="CENTER_BARRACRED"),
    "WEB": dict(aba_fonte="Planilha1", arquivo="WEB_BARRA_CRED.xlsx", aba_destino="WEB",
                sufixo="WEB_BARRACRED"),
}


def _ultimo_mes_pasta(pasta_canal):
    anos = sorted((p for p in os.listdir(pasta_canal) if p.isdigit()), reverse=True)
    for ano in anos:
        pasta_ano = os.path.join(pasta_canal, ano)
        meses = sorted((m for m in os.listdir(pasta_ano) if m.isdigit()), reverse=True)
        if meses:
            return os.path.join(pasta_ano, meses[0])
    raise FileNotFoundError(f"nenhuma pasta ano/mes em {pasta_canal}")


def _linhas_fonte(canal):
    cfg = CANAIS[canal]
    pasta_mes = _ultimo_mes_pasta(os.path.join(PASTA_FONTE_BASE, canal))
    arquivos = [f for f in os.listdir(pasta_mes)
                if f.lower().endswith(".xlsx") and f.upper().startswith(f"BARRA CRED {canal}")]
    if not arquivos:
        raise FileNotFoundError(f"nenhum AB 'BARRA CRED {canal} ...' em {pasta_mes}")
    caminho = os.path.join(pasta_mes, sorted(arquivos)[-1])

    wb = openpyxl.load_workbook(caminho, read_only=True, data_only=True)
    ws = wb[cfg["aba_fonte"]]
    linhas = list(ws.iter_rows(values_only=True))
    wb.close()

    header, *dados = linhas
    if len(header) < 16:
        raise ValueError(f"{caminho}: aba {cfg['aba_fonte']} tem {len(header)} colunas, esperava 16 (layout mudou?)")

    candidatas = [row for row in dados
                  if row[F_PEDIDO] is not None and row[F_STATUS_PEDIDO] == "Entregue"]
    return candidatas, caminho


def _data_prevista(txt):
    return datetime.datetime.strptime(txt, "%d/%m/%Y") if isinstance(txt, str) else txt


def _monta_center(row, seq):
    linha = [None] * 18
    linha[D_DATA_REPASSE] = _data_prevista(row[F_DT_PREVISTA])
    linha[D_MKT] = row[F_MKT]
    linha[D_SELLER] = row[F_SELLER]
    linha[D_DT_CRIACAO] = row[F_DT_CRIACAO]          # texto cru, igual a fonte
    linha[D_PEDIDO] = row[F_PEDIDO]
    linha[D_PEDIDO_SITE] = row[F_PEDIDO_SITE]
    linha[D_STATUS_PEDIDO] = row[F_STATUS_PEDIDO]
    linha[D_DT_STATUS_PEDIDO] = row[F_DT_STATUS_PEDIDO]  # texto cru, igual a fonte
    linha[D_TIPO] = row[F_TIPO]
    linha[D_STATUS_LANC] = row[F_STATUS_LANC]
    linha[D_DT_STATUS_LANC] = row[F_DT_STATUS_LANC]
    linha[D_DESCRICAO] = row[F_DESCRICAO]
    linha[D_VALOR] = row[F_VALOR]
    linha[D_DT_PREVISTA] = row[F_DT_PREVISTA]
    linha[D_ID_INTEGRADOR] = row[F_ID_INTEGRADOR]
    linha[D_PARCELA] = row[F_PARCELA]
    linha[D_SEQUENCIA] = seq
    linha[D_CHAVE] = f'=CONCATENATE(Q{{row}},"_","CENTER_BARRACRED")'
    return linha


def _monta_web(row, seq):
    linha = [None] * 18
    linha[D_DATA_REPASSE] = _data_prevista(row[F_DT_PREVISTA])
    linha[D_MKT] = row[F_MKT]
    linha[D_SELLER] = row[F_SELLER]
    linha[D_DT_CRIACAO] = _data_prevista(row[F_DT_CRIACAO])   # convertido (convencao do WEB)
    linha[D_PEDIDO] = row[F_PEDIDO]
    linha[D_PEDIDO_SITE] = None                               # convencao do WEB: sempre em branco
    linha[D_STATUS_PEDIDO] = row[F_STATUS_PEDIDO]
    dt_status = _data_prevista(row[F_DT_STATUS_PEDIDO])
    linha[D_DT_STATUS_PEDIDO] = to_excel(dt_status) if isinstance(dt_status, datetime.datetime) else dt_status
    linha[D_TIPO] = row[F_TIPO]
    linha[D_STATUS_LANC] = None                               # convencao do WEB: sempre em branco
    linha[D_DT_STATUS_LANC] = None                            # convencao do WEB: sempre em branco
    linha[D_DESCRICAO] = None                                 # convencao do WEB: sempre em branco
    linha[D_VALOR] = row[F_VALOR]
    linha[D_DT_PREVISTA] = row[F_DT_PREVISTA]
    linha[D_ID_INTEGRADOR] = row[F_ID_INTEGRADOR]
    linha[D_PARCELA] = row[F_PARCELA]
    linha[D_SEQUENCIA] = seq
    linha[D_CHAVE] = f'=CONCATENATE(Q{{row}},"_","WEB_BARRACRED")'
    return linha


_MONTA = {"CENTER": _monta_center, "WEB": _monta_web}
# campos que entram na chave natural de dedupe, por canal (WEB nao usa DT_STATUS_LANC:
# o destino desse campo e sempre None, nao discrimina nada)
_CHAVE_FONTE = {
    "CENTER": lambda row: (row[F_PEDIDO], row[F_TIPO], row[F_PARCELA], row[F_DT_STATUS_LANC], row[F_VALOR]),
    "WEB": lambda row: (row[F_PEDIDO], row[F_TIPO], row[F_PARCELA], row[F_VALOR]),
}
_CHAVE_DESTINO = {
    "CENTER": lambda r: (r[D_PEDIDO], r[D_TIPO], r[D_PARCELA], r[D_DT_STATUS_LANC], r[D_VALOR]),
    "WEB": lambda r: (r[D_PEDIDO], r[D_TIPO], r[D_PARCELA], r[D_VALOR]),
}


def _primeira_linha_vazia(ws, col_chave):
    ultima = 1
    for r in range(2, ws.max_row + 1):
        if ws.cell(row=r, column=col_chave).value is not None:
            ultima = r
    return ultima + 1


def empilhar(canal):
    cfg = CANAIS[canal]
    destino = os.path.join(PASTA_DESTINO, cfg["arquivo"])

    candidatas, caminho_fonte = _linhas_fonte(canal)

    wb = openpyxl.load_workbook(destino)
    ws = wb[cfg["aba_destino"]]

    linha_modelo = _primeira_linha_vazia(ws, D_PEDIDO + 1) - 1
    celulas_modelo = [ws.cell(row=linha_modelo, column=c) for c in range(1, 19)]
    ultima_seq = max((ws.cell(row=r, column=D_SEQUENCIA + 1).value or 0)
                      for r in range(2, linha_modelo + 1))

    existentes = {_CHAVE_DESTINO[canal]([ws.cell(row=r, column=c).value for c in range(1, 19)])
                  for r in range(2, linha_modelo + 1)}

    chave_fonte = _CHAVE_FONTE[canal]
    novas = [row for row in candidatas if chave_fonte(row) not in existentes]

    print(f"{canal}: fonte={os.path.basename(caminho_fonte)} | entregues={len(candidatas)} | novas={len(novas)}")
    if not novas:
        return

    os.makedirs(PASTA_BACKUP, exist_ok=True)
    backup = os.path.join(PASTA_BACKUP, f"{cfg['arquivo']}.bak_{datetime.datetime.now():%Y%m%d_%H%M%S}")
    shutil.copy2(destino, backup)
    print(f"  backup: {backup}")

    monta = _MONTA[canal]
    seq = ultima_seq
    proxima = linha_modelo + 1
    for i, row in enumerate(novas):
        r = proxima + i
        seq += 1
        linha = monta(row, seq)
        linha[D_CHAVE] = linha[D_CHAVE].format(row=r)
        for c, v in enumerate(linha, start=1):
            cel = ws.cell(row=r, column=c, value=v)
            modelo = celulas_modelo[c - 1]
            cel.number_format = modelo.number_format
            cel.font = copy(modelo.font)
            cel.alignment = copy(modelo.alignment)
            cel.fill = copy(modelo.fill)
            cel.border = copy(modelo.border)

    wb.save(destino)
    print(f"  gravado: {len(novas)} linhas novas, sequencia {ultima_seq + 1}..{seq}")


if __name__ == "__main__":
    if len(sys.argv) != 2 or sys.argv[1] not in CANAIS:
        sys.exit(f"uso: python empilhar_barra_cred.py <{'|'.join(CANAIS)}>")
    empilhar(sys.argv[1])
