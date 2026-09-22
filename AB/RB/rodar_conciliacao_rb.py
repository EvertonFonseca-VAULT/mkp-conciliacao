"""Roda a conciliacao com o RB: agrega o "Para Processar" de um canal/dia
por pedido, cruza com o extrato bancario, classifica o STATUS de cada
pedido (SEM DIFF / EST COM / REBATE / MEDIACAO / VERIFICAR) e -- pros
VERIFICAR -- busca no ML/MP os valores reais e sugere a acao + historico
de baixa. Substitui o ciclo manual AB_MERCADO_TOTVS.xlsm + CONTROLE_V2.xlsx
+ abrir ML pedido a pedido.

Uso (motor de classificacao, sem scraper -- roda com qualquer Python 3
com pandas/openpyxl):
    python rodar_conciliacao_rb.py <CANAL> <DATA>
    CANAL: FULL (unico canal ativo por enquanto -- ver canais.py)
    DATA: dd.mm ou dd.mm.aaaa, igual ao nome do arquivo "Para Processar" do dia

Uso com scraper do ML/MP (precisa do Chrome aberto com debug remoto, ja
logado -- ver Manuais/Documentação Screaper MP e ML.md -- e rodar com o
Python do venv do Screaper Mercado, que tem playwright):
    "...\\Screaper Mercado\\.venv\\Scripts\\python.exe" rodar_conciliacao_rb.py <CANAL> <DATA> --scraper
"""

import os
import shutil
import sys

import openpyxl
import pandas as pd

from agregador import agregar_por_pedido
from calendario_alvo import buscar_alvo
from canais import CAMINHO_RB_PADRAO, CANAIS, resolver_para_processar, validar_data_rb
from conferir_rbx import conferir as conferir_rbx, conferir_pedidos, rebate_lancado_por_pedido
from cruzamento import cruzar_com_rb
from descricoes_ml import gerar_descricoes_ml
from enriquecer import enriquecer
from parser_rb import carregar_rb, metadata_rb
from rebate_cmm import gerar_rebate_cmm
from rebate_processar import gerar_rebate_processar, salvar_rebate_processar

# Pasta onde o Everton arquiva uma copia do rbx.xlsx quando a Transitoria
# fica pronta pra transferencia (ver conferir_rbx.conferir -- chave
# "pronta_para_transferencia"). So um registro historico de qual RBX
# motivou qual transferencia -- NAO roda sozinho, ver _arquivar_rbx().
CAMINHO_ARQUIVO_MARKETPLACE = r"X:\109_GC_CTR_FINANCEIRO CONTAS A RECEBER\MKP\CONCILIAÇÕES DE MARKETPLACE\MERCADO"


def _arquivar_rbx(canal, data, raiz, caminho_rbx):
    """Copia o rbx.xlsx pra CAMINHO_ARQUIVO_MARKETPLACE/<ano>/<mes>.<ano>/
    <data>_<raiz>.xlsx e grava, na linha 3 (mesma linha de "Emissao:"), as
    formulas que o Everton usa pra provar visualmente que bateu: E3/F3 somam
    ENTRADAS/SAIDAS (SUBTOTAL, ignora linhas filtradas) e G3 = E3-F3 (o
    liquido). NAO chama sozinha em processar() -- so quando o Everton clica
    no botao "Arquivar RBX" do dashboard (gerado quando
    conferencia_rbx["pronta_para_transferencia"] fica True), pra nao
    arquivar antes da transferencia manual de verdade ter acontecido.
    Pedido do Everton, 17/09/2026."""
    dia, mes, ano = data.split(".")
    pasta_destino = os.path.join(CAMINHO_ARQUIVO_MARKETPLACE, ano, f"{mes}.{ano}")
    os.makedirs(pasta_destino, exist_ok=True)
    caminho_destino = os.path.join(pasta_destino, f"{data}_{raiz}.xlsx")
    shutil.copy2(caminho_rbx, caminho_destino)

    wb = openpyxl.load_workbook(caminho_destino)
    ws = wb.active
    ultima_linha = ws.max_row
    ws["E3"] = f"=SUBTOTAL(9,E6:E{ultima_linha})"
    ws["F3"] = f"=SUBTOTAL(9,F6:F{ultima_linha})"
    ws["G3"] = "=E3-F3"
    wb.save(caminho_destino)
    return caminho_destino


def rodar(caminho_para_processar, caminho_rb):
    para_processar = pd.read_excel(caminho_para_processar, sheet_name="Para Processar")
    rb = carregar_rb(caminho_rb)
    agregado = agregar_por_pedido(para_processar)
    return cruzar_com_rb(agregado, rb, para_processar), rb, para_processar


def _ordenar_fila_verificar(resultado):
    """Fila VERIFICAR em 2 grupos separados por sinal, nao por modulo -- o
    Everton lanca no TOTVS em lote (todo positivo primeiro, depois todo
    negativo, pra nao ficar abrindo o sistema toda hora): positivos do
    maior pro menor, negativos do mais negativo pro mais perto de zero.
    Ex.: 1000, 500, -1500, -300. Substitui a ordenacao por abs() (que
    misturava sinal e nao servia mais pra esse fluxo). Pedido do Everton,
    17/09/2026."""
    verificar = resultado[resultado["STATUS"] == "VERIFICAR"]
    positivos = verificar[verificar["DIF"] >= 0].sort_values("DIF", ascending=False)
    negativos = verificar[verificar["DIF"] < 0].sort_values("DIF", ascending=True)
    return pd.concat([positivos, negativos])


def _caminho_rebate_processar(pasta_processar_rebate, data_ref, rebate_processar):
    """None quando nao ha nada pra gerar (canal != FULL, ja filtrado antes
    de chamar isso -- rebate_processar vem None -- ou REBATE vazio no dia).
    Nome usa data_ref (dd/mm/aaaa, data real do RB) com pontos, mesma
    convencao dos arquivos que o processo automatico do TOTVS ja consome
    (ex. "PROCESSAR_MERCADO FULL dd.mm.aaaa.xlsx", ver canais.py) -- nao o
    `data` da linha de comando, que pode vir sem o ano (ex. "01.09")."""
    if rebate_processar is None or not len(rebate_processar):
        return None
    return os.path.join(pasta_processar_rebate, f"PROCESSAR_REBATE_{data_ref.replace('/', '.')}.xlsx")


def processar(canal, data, pasta_saida=".", usar_scraper=False):
    """Roda o ciclo inteiro (agregar/cruzar/enriquecer/rebate) e salva os
    xlsx. Usado tanto pelo `main()` (CLI) quanto pelo interativo.py
    (dashboard) -- unica funcao com a logica de ponta a ponta."""
    caminho_pp = resolver_para_processar(canal, data)
    caminho_rb = CAMINHO_RB_PADRAO
    raiz = CANAIS[canal]["raiz"]

    data_ref, responsavel = metadata_rb(caminho_rb)
    validar_data_rb(data, data_ref)

    resultado, rb, para_processar = rodar(caminho_pp, caminho_rb)
    verificar = resultado[resultado["STATUS"] == "VERIFICAR"]

    sufixo = f"{canal.lower()}_{data.replace('.', '')}"
    # "Descricoes", nao "Conciliacoes" -- desde 15/09/2026 o
    # descricoes_ml_*.xlsx salva em pasta_saida/Descricoes/ (ver mais
    # abaixo). Esse caminho aqui precisa bater com o de la, senao o cache
    # nunca acha o arquivo que ele mesmo salvou da ultima vez.
    caminho_descricoes_existente = os.path.join(pasta_saida, "Descricoes", f"descricoes_ml_{sufixo}.xlsx")

    ml_dados = None
    ml_de_cache = False
    if usar_scraper and len(verificar):
        from scraper_ml import resumir_pedidos  # so importa playwright se pedido

        print(f"Buscando {len(verificar)} pedido(s) VERIFICAR no ML/MP...")
        ml_dados = resumir_pedidos(list(verificar["PEDIDO_MKP"]))
    elif len(verificar):
        # Reaproveita uma raspagem anterior (descricoes_ml_<sufixo>.xlsx)
        # em vez de deixar a fila em branco -- so precisa rodar o scraper
        # de novo pros pedidos que ainda nao foram consultados. Pedido do
        # Everton, 16/09/2026.
        from classificar_ml import carregar_descricoes_salvas  # nao usa playwright

        ml_dados = carregar_descricoes_salvas(caminho_descricoes_existente)
        ml_de_cache = ml_dados is not None

    # Rebate ja movimentado na RBX hoje (credito normal ou cancelamento),
    # por pedido -- so pra AVISAR na Fila VERIFICAR quando o Rebate que o
    # scraper acha no ML/MP ja tem contrapartida no banco (ver
    # enriquecer.py). None se rbx.xlsx ainda nao existe -- nao trava o dia
    # a dia. Pedido do Everton, 22/09/2026.
    rebate_lancado_rbx = rebate_lancado_por_pedido(canal, os.path.join(pasta_saida, "AUXILIAR"))
    resultado = enriquecer(resultado, raiz, ml_dados, rebate_lancado_rbx)
    fila_verificar = _ordenar_fila_verificar(resultado)
    rebate_total = resultado.loc[resultado["STATUS"] == "REBATE", "REBATE"].sum()
    rebate_cmm = gerar_rebate_cmm(canal, resultado, raiz, data_ref, responsavel, rb)
    # PROCESSAR_REBATE: 2o layout do rebate, so pro FULL -- ver
    # rebate_processar.py. Nao substitui o Rebate CMM, e um output a mais
    # (importa direto no processo automatico do TOTVS, pasta
    # C:\RODAR TOTVS\PROCESSAR\REBATE\). Pedido do Everton, 17/09/2026.
    rebate_processar = gerar_rebate_processar(resultado, data_ref) if canal == "FULL" else None
    alvo = buscar_alvo(canal, data_ref)
    # Confere as 2 contas (original x transitoria) contra rb.xlsx/rbx.xlsx
    # em AUXILIAR/ -- None se ainda nao tem os 2 arquivos (feature opcional,
    # nao trava o dia a dia). Os 2 arquivos sao salvos manualmente pelo
    # Everton (pedido dele, 22/09/2026 -- sem auto-copia). Ver conferir_rbx.py.
    conferencia_rbx = conferir_rbx(canal, alvo, os.path.join(pasta_saida, "AUXILIAR"), rebate_total)
    # Fallback pedido a pedido (sem chave_rbx) pro canal que ainda nao tem
    # ela confirmada (ex. ES) -- so roda quando conferencia_rbx falhou por
    # isso (chave ausente), nao quando faltam rb.xlsx/rbx.xlsx (mesmo caso,
    # ja retorna None igual). Ver conferir_rbx.conferir_pedidos.
    valores_pedidos = (
        conferir_pedidos(os.path.join(pasta_saida, "AUXILIAR"), fila_verificar["PEDIDO_MKP"])
        if conferencia_rbx is None else None
    )

    # Descricoes irma de Conciliacoes/CMM (mesmo padrao ja usado na pasta do
    # Everton: AUXILIAR/CMM/Conciliacoes soltas na raiz) -- antes o
    # descricoes_ml_*.xlsx ia direto pra dentro de Conciliacoes/ e ficava
    # misturado com os conciliacao_*.xlsx. Pedido do Everton, 15/09/2026.
    pasta_conciliacoes = os.path.join(pasta_saida, "Conciliacoes")
    # "Rebate_CMM" (nao "CMM") -- o Everton ja usa "CMM" pra outra coisa (a
    # planilha CMM_MELI_FULL_Controle de Movimentos Manual de verdade), e o
    # nome curto colidia visualmente com essa pasta gerada pelo pipeline
    # (so tem o export auxiliar rebate_cmm_*.xlsx). Pedido do Everton,
    # 17/09/2026.
    pasta_cmm = os.path.join(pasta_saida, "Rebate_CMM")
    pasta_descricoes = os.path.join(pasta_saida, "Descricoes")
    # REBATE dentro de Conciliacoes/ (nao irma dela) -- pedido do Everton,
    # 17/09/2026.
    pasta_processar_rebate = os.path.join(pasta_conciliacoes, "REBATE")
    os.makedirs(pasta_conciliacoes, exist_ok=True)
    os.makedirs(pasta_cmm, exist_ok=True)
    os.makedirs(pasta_descricoes, exist_ok=True)
    caminho_xlsx = os.path.join(pasta_conciliacoes, f"conciliacao_{sufixo}.xlsx")
    caminho_rebate_xlsx = os.path.join(pasta_cmm, f"rebate_cmm_{sufixo}.xlsx")
    caminho_rebate_processar_xlsx = _caminho_rebate_processar(pasta_processar_rebate, data_ref, rebate_processar)
    if caminho_rebate_processar_xlsx:
        os.makedirs(pasta_processar_rebate, exist_ok=True)

    descricoes_ml = gerar_descricoes_ml(ml_dados)
    caminho_descricoes_xlsx = (
        os.path.join(pasta_descricoes, f"descricoes_ml_{sufixo}.xlsx")
        if len(descricoes_ml) else None
    )

    # AB (Para Processar) importado junto -- so pra nao precisar abrir o
    # arquivo original a parte enquanto o time se acostuma com essa planilha.
    # Remover quando nao for mais necessario.
    try:
        with pd.ExcelWriter(caminho_xlsx) as writer:
            resultado.to_excel(writer, sheet_name="Conciliacao", index=False)
            fila_verificar.to_excel(writer, sheet_name="Fila VERIFICAR", index=False)
            rebate_cmm.to_excel(writer, sheet_name="Rebate CMM", index=False)
            para_processar.to_excel(writer, sheet_name="AB", index=False)

        with pd.ExcelWriter(caminho_rebate_xlsx) as writer:
            rebate_cmm.to_excel(writer, sheet_name="Rebate CMM", index=False)

        if caminho_descricoes_xlsx:
            descricoes_ml.to_excel(caminho_descricoes_xlsx, index=False)

        if caminho_rebate_processar_xlsx:
            salvar_rebate_processar(rebate_processar, caminho_rebate_processar_xlsx, data_ref)
    except PermissionError as exc:
        raise PermissionError(
            f"Feche o arquivo no Excel antes de rodar de novo: {exc.filename}"
        ) from exc

    return {
        "resultado": resultado,
        "fila_verificar": fila_verificar,
        "rebate_cmm": rebate_cmm,
        "rebate_processar": rebate_processar,
        "alvo": alvo,
        "ml_atualizado": ml_dados is not None and not ml_de_cache,
        "ml_de_cache": ml_de_cache,
        "rb": rb,
        "para_processar": para_processar,
        "data_ref": data_ref,
        "responsavel": responsavel,
        "raiz": raiz,
        "caminho_pp": caminho_pp,
        "caminho_rb": caminho_rb,
        "caminho_xlsx": caminho_xlsx,
        "caminho_rebate_xlsx": caminho_rebate_xlsx,
        "caminho_rebate_processar_xlsx": caminho_rebate_processar_xlsx,
        "caminho_descricoes_xlsx": caminho_descricoes_xlsx,
        "conferencia_rbx": conferencia_rbx,
        "valores_pedidos": valores_pedidos,
    }


def main():
    usar_scraper = "--scraper" in sys.argv
    args = [a for a in sys.argv[1:] if a != "--scraper"]

    canal = args[0].upper()
    if canal not in CANAIS:
        sys.exit(f"Canal '{canal}' desconhecido. Opcoes: {', '.join(CANAIS)}")
    data = args[1]

    try:
        r = processar(canal, data, usar_scraper=usar_scraper)
    except ValueError as exc:
        sys.exit(str(exc))

    resultado, fila_verificar, rebate_cmm = r["resultado"], r["fila_verificar"], r["rebate_cmm"]

    print(f"Para Processar: {r['caminho_pp']}")
    print(f"RB: {r['caminho_rb']}")
    print(f"Total de pedidos: {len(resultado)}")
    for status, qtd in resultado["STATUS"].value_counts().items():
        print(f"  {status}: {qtd}")
    print(f"Fila VERIFICAR: {len(fila_verificar)}")
    print(f"Rebate CMM (linhas pra colar no Movimento_Bancario): {len(rebate_cmm)}")
    print(f"Salvo em: {os.path.abspath(r['caminho_xlsx'])}")


def _testar_arquivar_rbx():
    """_arquivar_rbx nao roda sozinha em processar() -- so quando chamada.
    Confere caminho de destino (<ano>/<mes>.<ano>/<data>_<raiz>.xlsx) e as
    formulas E3/F3/G3. Pedido do Everton, 17/09/2026."""
    import tempfile

    global CAMINHO_ARQUIVO_MARKETPLACE

    with tempfile.TemporaryDirectory() as tmp:
        caminho_rbx = os.path.join(tmp, "rbx.xlsx")
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.append(["Dt.Ref: 01/09/2026"])
        ws.append(["Hora: 10:00:00"])
        ws.append(["Emissão: 01/09/2026"])
        ws.append(["BANCO:", "237 - BRADESCO SA", "AGENCIA:", "3708", "CONTA:", "X103925", "SALDO INICIAL:", "0,00"])
        ws.append(["DATA", "OPERACAO", "DOCUMENTO", "PREFIXO/TITULO", "ENTRADAS", "SAIDAS",
                    "SALDO ATUAL", "USUARIO", "PED MKP", "Vendedor", "Emissao Titulo", "Saldo Titulo"])
        ws.append(["01/09/2026", "REC MKP MKP_MELI_FU REF PARC 1", "1", "", 100.0, 0, 0, "", "1", "", "", 0])
        ws.append(["01/09/2026", "MKP_FULL_REPASSE", "", "", 0, 100.0, 0, "", "", "", "", 0])
        wb.save(caminho_rbx)

        anterior = CAMINHO_ARQUIVO_MARKETPLACE
        CAMINHO_ARQUIVO_MARKETPLACE = tmp
        try:
            caminho_destino = _arquivar_rbx("FULL", "01.09.2026", "MKP_MELI_FU", caminho_rbx)
        finally:
            CAMINHO_ARQUIVO_MARKETPLACE = anterior

        assert caminho_destino == os.path.join(tmp, "2026", "09.2026", "01.09.2026_MKP_MELI_FU.xlsx"), caminho_destino
        assert os.path.exists(caminho_destino)

        wb2 = openpyxl.load_workbook(caminho_destino)
        ws2 = wb2.active
        assert ws2["E3"].value == "=SUBTOTAL(9,E6:E7)", ws2["E3"].value
        assert ws2["F3"].value == "=SUBTOTAL(9,F6:F7)", ws2["F3"].value
        assert ws2["G3"].value == "=E3-F3", ws2["G3"].value


def _testar_ordenar_fila_verificar():
    """Exemplo real do Everton: DIF = [1000, 500, -1500, -300] tem que sair
    1000, 500, -1500, -300 (positivos maior->menor, depois negativos
    mais-negativo->mais-perto-de-zero) -- lote de lancamento no TOTVS."""
    resultado = pd.DataFrame({
        "PEDIDO_MKP": ["a", "b", "c", "d", "e"],
        "STATUS": ["VERIFICAR", "VERIFICAR", "VERIFICAR", "VERIFICAR", "SEM DIFF"],
        "DIF": [500.0, 1000.0, -300.0, -1500.0, 999.0],
    })
    fila = _ordenar_fila_verificar(resultado)
    assert list(fila["DIF"]) == [1000.0, 500.0, -1500.0, -300.0], list(fila["DIF"])


def _testar_caminho_rebate_processar():
    """So gera arquivo/pasta quando tem REBATE de verdade pra colocar
    dentro -- canal != FULL ja chega aqui como rebate_processar=None (ver
    processar())."""
    rebate_vazio = pd.DataFrame({"NR PEDIDO MKP": []})
    rebate_com_linha = pd.DataFrame({"NR PEDIDO MKP": ["1"]})

    pasta = r"C:\saida\Conciliacoes\REBATE"
    assert _caminho_rebate_processar(pasta, "02/09/2026", None) is None
    assert _caminho_rebate_processar(pasta, "02/09/2026", rebate_vazio) is None
    assert _caminho_rebate_processar(pasta, "02/09/2026", rebate_com_linha) == (
        r"C:\saida\Conciliacoes\REBATE\PROCESSAR_REBATE_02.09.2026.xlsx"
    )


if __name__ == "__main__":
    if "--self-test" in sys.argv:
        _testar_arquivar_rbx()
        _testar_ordenar_fila_verificar()
        _testar_caminho_rebate_processar()
        print("rodar_conciliacao_rb.py: ok")
    else:
        main()
