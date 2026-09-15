"""Parser do extrato bancario RB (CONCILIADOR_RB.xlsx), exportado paginado
do TOTVS. O relatorio repete o bloco de cabecalho (BANCO/AGENCIA/CONTA +
linha de colunas) a cada ~57 linhas (uma pagina do relatorio original) e
termina com um bloco de totais. Ver documentacao em
Manuais/Passo a passo - Conciliacao com RB (CONTROLE_V2).md.
"""

import openpyxl
import pandas as pd

COLUNAS_LANCAMENTO = [
    "DATA", "OPERACAO", "DOCUMENTO", "PREFIXO_TITULO", "ENTRADAS", "SAIDAS",
    "SALDO_ATUAL", "USUARIO", "PED_MKP", "VENDEDOR", "EMISSAO_TITULO", "SALDO_TITULO",
]


def metadata_rb(caminho):
    """Data de referencia do RB ('Dt.Ref: dd/mm/aaaa' em A1) e primeiro nome
    de quem salvou o arquivo (metadata do proprio Excel) -- usado pra
    validar que o RB e do dia certo antes de rodar, e pra preencher a
    aba Rebate CMM sem reabrir o arquivo depois."""
    wb = openpyxl.load_workbook(caminho, data_only=True)
    texto = wb.active.cell(row=1, column=1).value or ""
    data_ref = texto.split(":", 1)[1].strip() if ":" in texto else None
    quem = (wb.properties.lastModifiedBy or wb.properties.creator or "").strip()
    return data_ref, quem.split(" ")[0].upper() if quem else ""


def carregar_rb(caminho):
    wb = openpyxl.load_workbook(caminho, data_only=True)
    ws = wb[wb.sheetnames[0]]

    linhas = []
    banco_atual = agencia_atual = conta_atual = None

    for row in ws.iter_rows(min_row=1, values_only=True):
        primeira, segunda = row[0], row[1]

        if primeira == "BANCO:":
            banco_atual, agencia_atual, conta_atual = row[1], row[3], row[5]
            continue
        if primeira == "DATA" and segunda == "OPERACAO":
            continue  # cabecalho de colunas, repete a cada pagina
        if segunda == "NAO CONCILIADOS":
            break  # bloco de totais no fim do relatorio
        if segunda is None:
            continue  # metadado (Dt.Ref/Hora/Emissao) ou linha vazia

        linhas.append(list(row[:12]) + [banco_atual, agencia_atual, conta_atual])

    df = pd.DataFrame(linhas, columns=COLUNAS_LANCAMENTO + ["BANCO", "AGENCIA", "CONTA"])
    for col in ("ENTRADAS", "SAIDAS", "SALDO_ATUAL", "SALDO_TITULO"):
        df[col] = pd.to_numeric(df[col], errors="coerce").fillna(0.0)
    df["PED_MKP"] = df["PED_MKP"].apply(lambda v: "" if v is None else str(v).strip())
    return df
