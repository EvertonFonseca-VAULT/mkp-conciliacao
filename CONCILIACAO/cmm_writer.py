"""Inclui linhas na aba Movimento_Bancario do CMM real -- mesmo jeito que
EMPILHAR/empilhar_kabum.py (e empilhar_barra_cred.py) ja gravam nos arquivos
"empilhados": openpyxl puro (load_workbook + editar + save), com BACKUP
automatico do arquivo original antes de salvar, igual PASTA_BACKUP de la.

Historico desta sessao (pra quem ler depois): a 1a versao usava Excel de
verdade via COM (win32com) pra nao perder a Data Validation "extension"
(dropdown x14) que o openpyxl avisa que vai remover em QUALQUER
abrir+salvar desse arquivo. Na pratica o COM deu muito mais problema (Excel
preso rejeitando chamada, processo orfao, recalculo de 32 mil linhas
travando o Quit) do que vale a pena -- e o pipeline JA aceita esse mesmo
risco pros arquivos empilhados (sem reclamacao conhecida). Trocado pra
openpyxl + backup, igual o resto do projeto. Pedido do Everton, 14/10/2026.

NATUREZA (col G) e HISTORICO (col M): no CMM do FULL sao FORMULA (XLOOKUP e
CONCATENATE) -- copiamos a formula da ultima linha pra linha nova, ajustando
so o numero da linha (igual arrastar no Excel). MAS isso NAO e universal:
testado contra o CMM real do WEB (14/10/2026) e a coluna G la e VALOR FIXO
(1107 escrito direto, sem formula) -- layout difere por arquivo. Por isso
_valor_ou_formula checa a ULTIMA linha de cada arquivo e decide na hora: se
for formula, copia formula; se for valor, escreve o literal que o chamador
mandou (rebate_cmm.py e o dashboard JA calculam NATUREZA/HISTORICO certos
mesmo sem formula -- so nao usavamos esse valor antes).
"""

import datetime
import os
import re
import shutil

import openpyxl

COLUNA_LETRA = {
    "RESPONSAVEL": "A", "DATA_LACTO": "B", "ORIGEM": "C", "PAGADOR": "D", "LOJA": "E",
    "MOTIVO DO MOVIMENTO": "F", "PEDIDO MKP / CV": "H", "VENDEDOR": "I",
    "BANCO": "J", "AGENCIA": "K", "CONTA": "L", "VALOR R$": "N",
    "OBSERVACAO": "P", "PEDIDO_ERP": "Q", "PEDIDO_SITE": "R", "N. FISCAL": "S", "PROCESSO": "T",
}
# NATUREZA/HISTORICO: coluna pode ser formula OU valor fixo, depende do
# arquivo (ver docstring) -- resolvido na hora, nao fixo como COLUNA_LETRA.
COLUNA_LETRA_NATUREZA_HISTORICO = {"NATUREZA": "G", "HISTORICO": "M"}
_PREFIXO_ABA = "Movimento_Banc"  # nome real tem acento (Movimento_Bancário)
PASTA_BACKUP = r"C:\BACKUP CMM"  # mesmo padrao de EMPILHAR\empilhar_kabum.PASTA_BACKUP


def _nome_aba_real(wb):
    for nome in wb.sheetnames:
        if nome.startswith(_PREFIXO_ABA):
            return nome
    raise ValueError("Aba Movimento_Bancario nao encontrada no arquivo.")


def _copiar_formula_para_linha(ws, coluna, linha_origem, linha_destino):
    formula = ws[f"{coluna}{linha_origem}"].value
    nova_formula = re.sub(rf"(?<!\d){linha_origem}(?!\d)", str(linha_destino), formula)
    ws[f"{coluna}{linha_destino}"] = nova_formula


def _escrever_natureza_historico(ws, campo, letra, linha_origem, linha_destino, linha):
    """Se a ultima linha dessa coluna e formula, copia (arrastar). Senao, e
    valor fixo nesse arquivo (ex.: WEB) -- escreve o literal que o chamador
    ja calculou (rebate_cmm.py / dashboard.js ja mandam NATUREZA/HISTORICO
    certos mesmo sem formula)."""
    referencia = ws[f"{letra}{linha_origem}"].value
    if isinstance(referencia, str) and referencia.startswith("="):
        _copiar_formula_para_linha(ws, letra, linha_origem, linha_destino)
        return
    valor = linha.get(campo)
    if valor not in (None, ""):
        ws[f"{letra}{linha_destino}"] = valor


def _fazer_backup(caminho):
    os.makedirs(PASTA_BACKUP, exist_ok=True)
    nome = os.path.basename(caminho)
    destino = os.path.join(PASTA_BACKUP, f"{nome}.bak_{datetime.datetime.now():%Y%m%d_%H%M%S}")
    shutil.copy2(caminho, destino)
    return destino


def incluir_linhas_cmm(caminho_xlsx, linhas):
    """Faz backup do CMM, acha a 1a linha vazia da aba Movimento_Bancario,
    escreve `linhas` (dict com as chaves de COLUNA_LETRA) e salva. Devolve
    (ultima_linha, caminho_do_backup)."""
    if not linhas:
        return None, None
    backup = _fazer_backup(caminho_xlsx)
    wb = openpyxl.load_workbook(caminho_xlsx)
    ws = wb[_nome_aba_real(wb)]
    # NAO usa ws.max_row direto -- conta linhas em branco so formatadas
    # (confirmado: no WEB real max_row veio 1048576). Mesmo jeito de
    # EMPILHAR/empilhar_kabum._primeira_linha_vazia: varre a coluna A.
    ultima_linha = 3  # cabecalho
    for r in range(4, ws.max_row + 1):
        if ws.cell(row=r, column=1).value is not None:
            ultima_linha = r
    if ultima_linha < 4:
        raise ValueError("Nao achei nenhuma linha de dado (cabecalho acima da linha 4?).")

    for i, linha in enumerate(linhas):
        destino = ultima_linha + 1 + i
        for campo, letra in COLUNA_LETRA.items():
            valor = linha.get(campo)
            if valor not in (None, ""):
                ws[f"{letra}{destino}"] = valor
        for campo, letra in COLUNA_LETRA_NATUREZA_HISTORICO.items():
            _escrever_natureza_historico(ws, campo, letra, ultima_linha, destino, linha)

    wb.save(caminho_xlsx)
    return ultima_linha + len(linhas), backup


def linhas_do_clipboard_cmm(texto, colunas_ordem):
    """Converte o texto tab-separado que o botao 'Copiar' do dashboard
    ja gera (mesma ordem de rebate_cmm.COLUNAS) numa lista de dict pronta
    pra incluir_linhas_cmm. Linha em branco e ignorada."""
    linhas = []
    for linha_txt in texto.splitlines():
        if not linha_txt.strip():
            continue
        campos = linha_txt.split("\t")
        linhas.append(dict(zip(colunas_ordem, campos)))
    return linhas


if __name__ == "__main__":
    import sys
    import tempfile

    if "--self-test-real-excel" not in sys.argv:
        print("Self-test precisa de um arquivo CMM real de referencia (layout) -- rode com "
              "--self-test-real-excel contra uma COPIA de teste (nunca o arquivo de rede de verdade).")
        sys.exit(0)

    caminho_teste = sys.argv[sys.argv.index("--self-test-real-excel") + 1]
    with tempfile.TemporaryDirectory() as tmp:
        copia = os.path.join(tmp, os.path.basename(caminho_teste))
        shutil.copy(caminho_teste, copia)

        linhas = [{
            "RESPONSAVEL": "TESTE_PROTOTIPO", "DATA_LACTO": "01/01/2026", "ORIGEM": "MKP",
            "PAGADOR": "MELI", "LOJA": "FULL", "MOTIVO DO MOVIMENTO": "RECEBIMENTO DE REBATE",
            "NATUREZA": 1107, "HISTORICO": "MKP_TESTE_RECEBIMENTO DE REBATE",  # usado SO se a
            # coluna real for valor fixo (ver _escrever_natureza_historico) -- se for formula,
            # esses 2 campos sao ignorados (a formula copiada que calcula).
            "PEDIDO MKP / CV": "999999999", "VENDEDOR": "751", "BANCO": "237",
            "AGENCIA": "3708", "CONTA": "X103925", "VALOR R$": 1.23, "OBSERVACAO": "",
        }]
        ultima, backup = incluir_linhas_cmm(copia, linhas)
        print(f"Escreveu na linha {ultima}; backup em {backup}")
        assert os.path.exists(backup)

        wb = openpyxl.load_workbook(copia, data_only=False)
        ws = wb[_nome_aba_real(wb)]
        linha_escrita = [c.value for c in ws[ultima]]
        assert linha_escrita[0] == "TESTE_PROTOTIPO", linha_escrita
        assert linha_escrita[13] == 1.23, linha_escrita  # VALOR R$ (col N, indice 13)
        natureza, historico = linha_escrita[6], linha_escrita[12]
        eh_formula = lambda v: isinstance(v, str) and v.startswith("=")
        print(f"NATUREZA: {'formula' if eh_formula(natureza) else 'valor fixo'} ({natureza!r})")
        print(f"HISTORICO: {'formula' if eh_formula(historico) else 'valor fixo'} ({historico!r})")
        assert natureza not in (None, ""), "NATUREZA ficou vazia"
        assert historico not in (None, ""), "HISTORICO ficou vazio"
        if not eh_formula(natureza):
            assert natureza == 1107, natureza  # veio do literal passado em `linhas`
        os.remove(backup)
        print("cmm_writer.py: ok (formula OU valor fixo conforme o arquivo, valor gravado, backup criado)")
