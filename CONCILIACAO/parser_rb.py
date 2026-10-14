"""Parser do extrato bancario RB (CONCILIADOR_RB.xlsx), exportado paginado
do TOTVS. O relatorio repete o bloco de cabecalho (BANCO/AGENCIA/CONTA +
linha de colunas) a cada ~57 linhas (uma pagina do relatorio original) e
termina com um bloco de totais. Ver documentacao em
Manuais/Passo a passo - Conciliacao com RB (CONTROLE_V2).md.
"""

import openpyxl
import pandas as pd

from ids import limpar_id

COLUNAS_LANCAMENTO = [
    "DATA", "OPERACAO", "DOCUMENTO", "PREFIXO_TITULO", "ENTRADAS", "SAIDAS",
    "SALDO_ATUAL", "USUARIO", "PED_MKP", "VENDEDOR", "EMISSAO_TITULO", "SALDO_TITULO",
]


def metadata_rb(caminho):
    """Data de referencia do RB ('Dt.Ref: dd/mm/aaaa' em A1) e primeiro nome
    de quem salvou o arquivo (metadata do proprio Excel) -- usado pra
    validar que o RB e do dia certo antes de rodar, e pra preencher a
    aba Rebate CMM sem reabrir o arquivo depois."""
    wb = openpyxl.load_workbook(caminho, data_only=True)
    texto = wb.active.cell(row=1, column=1).value or ""
    data_ref = texto.split(":", 1)[1].strip() if ":" in texto else None
    quem = (wb.properties.lastModifiedBy or wb.properties.creator or "").strip()
    return data_ref, quem.split(" ")[0].upper() if quem else ""


def carregar_rb(caminho):
    wb = openpyxl.load_workbook(caminho, data_only=True)
    ws = wb[wb.sheetnames[0]]

    linhas = []
    banco_atual = agencia_atual = conta_atual = None

    for row in ws.iter_rows(min_row=1, values_only=True):
        primeira, segunda = row[0], row[1]

        if primeira == "BANCO:":
            banco_atual, agencia_atual, conta_atual = row[1], row[3], row[5]
            continue
        if primeira == "DATA" and segunda == "OPERACAO":
            continue  # cabecalho de colunas, repete a cada pagina
        if segunda == "NAO CONCILIADOS":
            break  # bloco de totais no fim do relatorio
        if segunda is None:
            continue  # metadado (Dt.Ref/Hora/Emissao) ou linha vazia

        linhas.append(list(row[:12]) + [banco_atual, agencia_atual, conta_atual])

    df = pd.DataFrame(linhas, columns=COLUNAS_LANCAMENTO + ["BANCO", "AGENCIA", "CONTA"])
    for col in ("ENTRADAS", "SAIDAS", "SALDO_ATUAL", "SALDO_TITULO"):
        df[col] = pd.to_numeric(df[col], errors="coerce").fillna(0.0)
    # limpar_id (mesma normalizacao do PEDIDO_MKP em agregador.py), nao um
    # str(v).strip() cru: se a celula PED MKP do RB/RBX for lida como numero
    # pelo openpyxl, "2000017834651356.0" nunca bate com o PEDIDO_MKP
    # correspondente (que ja passa por limpar_id e fica sem o ".0") --
    # achado com o pedido 2000017834651356 da linha 1490, que tinha debito
    # lancado no RB mas aparecia como "nao encontrado" no dashboard.
    # limpar_id tambem cobre celula vazia -> "" (pandas converte pra NaN,
    # nao None, ao montar a coluna -- "v is None" sozinho deixava passar
    # como a string literal "nan", que vazava pro rebate_cmm.py e pro
    # groupby do cruzamento.py). Pedido do Everton, 15/09/2026.
    df["PED_MKP"] = df["PED_MKP"].apply(limpar_id)
    return df


if __name__ == "__main__":
    import tempfile
    import os

    def _escrever(caminho, linhas):
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.append(["Dt.Ref: 01/09/2026"])
        ws.append(["BANCO:", "237 - BRADESCO SA", "AGENCIA:", "3708", "CONTA:", "103925", "SALDO INICIAL:", "0,00"])
        ws.append(["DATA", "OPERACAO", "DOCUMENTO", "PREFIXO/TITULO", "ENTRADAS", "SAIDAS",
                    "SALDO ATUAL", "USUARIO", "PED MKP", "Vendedor", "Emissao Titulo", "Saldo Titulo"])
        for operacao, pedido in linhas:
            ws.append(["01/09/2026", operacao, pedido, "", 1, 0, 0, "", pedido, "", "", 0])
        ws.append(["", "NAO CONCILIADOS", "CONCILIADOS", "TOTAL"])
        wb.save(caminho)

    with tempfile.TemporaryDirectory() as tmp:
        caminho = os.path.join(tmp, "rb.xlsx")
        _escrever(caminho, [
            ("MKP_FULL_REPASSE", ""),
            ("CAN-REC MKP REF PARC 1", "1"),
            ("CAN-REC MKP REF PARC 2", 2000017834651356),  # celula numerica, nao texto
        ])
        df = carregar_rb(caminho)
        # PED_MKP em branco tem que virar "" -- nao a string "nan" (pandas
        # converte celula vazia pra NaN, nao pra None, ao montar a coluna).
        # E celula numerica nao pode sobrar ".0" no fim (senao nunca bate
        # com o PEDIDO_MKP do lado do agregador, que ja usa limpar_id).
        assert df["PED_MKP"].tolist() == ["", "1", "2000017834651356"], df["PED_MKP"].tolist()
        print("parser_rb.py: ok")
