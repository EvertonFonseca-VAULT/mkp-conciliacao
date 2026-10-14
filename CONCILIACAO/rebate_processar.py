"""Gera o layout PROCESSAR_REBATE (so canal FULL) -- 2o output do rebate,
ALEM do XLSX Rebate CMM que ja existe (rebate_cmm.py, nao mexe nele). Esse
aqui e o layout que o processo automatico de import do TOTVS espera (pasta
C:\\RODAR TOTVS\\PROCESSAR\\REBATE\\), estrutura conferida direto num arquivo
real de la (PROCESSAR_REBATE_02.09.2026.xlsx). Pedido do Everton, 17/09/2026.
"""

import datetime

import pandas as pd

COLUNAS = ["NR PEDIDO MKP", "VALOR REPASSE", "DATA", "TIPO", "COMISSAO"]


def _data_ref_para_date(data_ref):
    """data_ref vem de metadata_rb() no formato dd/mm/aaaa -- a celula B1/
    coluna DATA no arquivo real e datetime de verdade, nao texto."""
    return datetime.datetime.strptime(data_ref, "%d/%m/%Y").date()


def gerar_rebate_processar(resultado, data_ref):
    data = _data_ref_para_date(data_ref)
    rebates = resultado[resultado["STATUS"] == "REBATE"]
    linhas = [{
        "NR PEDIDO MKP": row["PEDIDO_MKP"],
        "VALOR REPASSE": row["REBATE"],
        "DATA": data,
        "TIPO": "Venda" if row["REBATE"] >= 0 else "Estorno",
        "COMISSAO": 0,
    } for _, row in rebates.iterrows()]
    return pd.DataFrame(linhas, columns=COLUNAS)


def salvar_rebate_processar(df, caminho_xlsx, data_ref):
    """Layout fixo (conferido no arquivo real): sheet unica "MKP_AB", linha 1
    mistura metadado (A1:C1) e cabecalho dos dados (D1:H1), dados a partir de
    D2. startcol=3 poe o DataFrame (com header) comecando na coluna D.
    B1 e a coluna F (DATA) usam number_format "mm-dd-yy" igual ao arquivo
    real -- sem isso o Excel mostra a data como ano-mes-dia. Everton notou
    isso no primeiro arquivo gerado, 17/09/2026."""
    with pd.ExcelWriter(caminho_xlsx, engine="openpyxl") as writer:
        df.to_excel(writer, sheet_name="MKP_AB", index=False, startcol=3)
        ws = writer.sheets["MKP_AB"]
        ws["A1"] = "DATA REPASSE"
        ws["B1"] = _data_ref_para_date(data_ref)
        ws["B1"].number_format = "mm-dd-yy"
        ws["C1"] = "MERCADO FULL"
        for linha in range(2, len(df) + 2):
            ws.cell(row=linha, column=6).number_format = "mm-dd-yy"  # coluna F = DATA


if __name__ == "__main__":
    import os
    import tempfile

    import openpyxl

    resultado = pd.DataFrame({
        "PEDIDO_MKP": ["1", "2", "3"],
        "STATUS": ["REBATE", "REBATE", "VERIFICAR"],
        "REBATE": [30.0, -30.5, 999.0],
    })
    df = gerar_rebate_processar(resultado, "02/09/2026")
    assert list(df.columns) == COLUNAS
    assert len(df) == 2  # so as REBATE, a VERIFICAR fica de fora
    assert list(df["TIPO"]) == ["Venda", "Estorno"]
    assert (df["COMISSAO"] == 0).all()
    assert df["DATA"].iloc[0] == datetime.date(2026, 9, 2)

    with tempfile.TemporaryDirectory() as tmp:
        caminho = os.path.join(tmp, "PROCESSAR_REBATE_02.09.2026.xlsx")
        salvar_rebate_processar(df, caminho, "02/09/2026")

        wb = openpyxl.load_workbook(caminho)
        ws = wb["MKP_AB"]
        assert ws["A1"].value == "DATA REPASSE"
        assert ws["B1"].value == datetime.datetime(2026, 9, 2)  # openpyxl le date como datetime
        assert ws["B1"].number_format == "mm-dd-yy"
        assert ws["F2"].number_format == "mm-dd-yy" and ws["F3"].number_format == "mm-dd-yy"
        assert ws["C1"].value == "MERCADO FULL"
        assert ws["D1"].value == "NR PEDIDO MKP"
        assert ws["E1"].value == "VALOR REPASSE"
        assert ws["F1"].value == "DATA"
        assert ws["G1"].value == "TIPO"
        assert ws["H1"].value == "COMISSAO"
        assert ws["D2"].value == "1"
        assert ws["E2"].value == 30.0
        assert ws["G2"].value == "Venda"
        assert ws["D3"].value == "2"
        assert ws["G3"].value == "Estorno"
        assert ws["H2"].value == 0 and ws["H3"].value == 0

    # 41 pedidos reais do dia 02/09/2026 (lidos do PROCESSAR_REBATE_02.09.2026.xlsx
    # de producao pelo Everton) -- confere linha a linha que TIPO/COMISSAO batem,
    # em particular o unico Estorno do lote (pedido 2000017834924604, -30,5).
    _REAIS_02_09 = [
        ("2000017834924604", -30.5, "Estorno"), ("2000017771286522", 79.79, "Venda"),
        ("2000017887346980", 70, "Venda"), ("2000017904668762", 22.5, "Venda"),
        ("2000017944118140", 57.7, "Venda"), ("2000017957155784", 36.72, "Venda"),
        ("2000017999272824", 15.04, "Venda"), ("2000018001112988", 75.84, "Venda"),
        ("2000018006603830", 60, "Venda"), ("2000018025209720", 119, "Venda"),
        ("2000018025385016", 119, "Venda"), ("2000018037510094", 59.5, "Venda"),
        ("2000018037707128", 4.8, "Venda"), ("2000018044841130", 96.61, "Venda"),
        ("2000018046091186", 20, "Venda"), ("2000018047595256", 96.61, "Venda"),
        ("2000018047675078", 128.6, "Venda"), ("2000018049402648", 30, "Venda"),
        ("2000018051194114", 60, "Venda"), ("2000018052201548", 119, "Venda"),
        ("2000018054894360", 60, "Venda"), ("2000018055751674", 9, "Venda"),
        ("2000018056354306", 119, "Venda"), ("2000018057429522", 57.7, "Venda"),
        ("2000018063233868", 87, "Venda"), ("2000018064430680", 119, "Venda"),
        ("2000018064779946", 20, "Venda"), ("2000018064940554", 49.8, "Venda"),
        ("2000018065710910", 135, "Venda"), ("2000018066080440", 47.6, "Venda"),
        ("2000018066277922", 96.61, "Venda"), ("2000018066371054", 70, "Venda"),
        ("2000018068316020", 87, "Venda"), ("2000018069677766", 96.61, "Venda"),
        ("2000018073921508", 20, "Venda"), ("2000018075512364", 109.8, "Venda"),
        ("2000018079160088", 87, "Venda"), ("2000018080390946", 87, "Venda"),
        ("2000018080703758", 139.44, "Venda"), ("2000018083360946", 166.6, "Venda"),
        ("2000018086650728", 57.7, "Venda"),
    ]
    resultado_real = pd.DataFrame({
        "PEDIDO_MKP": [p for p, _, _ in _REAIS_02_09],
        "STATUS": ["REBATE"] * len(_REAIS_02_09),
        "REBATE": [v for _, v, _ in _REAIS_02_09],
    })
    df_real = gerar_rebate_processar(resultado_real, "02/09/2026")
    assert len(df_real) == 41
    assert list(df_real["TIPO"]) == [t for _, _, t in _REAIS_02_09]
    assert list(df_real["NR PEDIDO MKP"]) == [p for p, _, _ in _REAIS_02_09]
    assert (df_real["COMISSAO"] == 0).all()
    assert sum(df_real["TIPO"] == "Estorno") == 1

    print("rebate_processar.py: ok")
