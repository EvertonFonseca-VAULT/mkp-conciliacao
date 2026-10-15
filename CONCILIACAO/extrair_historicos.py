"""Extrai a aba HISTORICOS do CONTROLE_V2.xlsx (loja -> RAIZ do historico
TOTVS) pra historicos_totvs.json, salvo no projeto -- assim dashboard.py
nao depende de abrir o CONTROLE_V2.xlsx (arquivo local, so existe na
maquina de quem tem o Conciliador Backup) toda vez que mostra o painel de
referencia. Reroda isso quando a planilha mudar (loja nova, raiz trocada).

Uso: python extrair_historicos.py
"""

import json
import os

import openpyxl

CAMINHO_CONTROLE_V2 = r"C:\TEMP\Conciliador - Backup\CONTROLE_V2.xlsx"
CAMINHO_SAIDA = os.path.join(os.path.dirname(__file__), "historicos_totvs.json")


def extrair():
    wb = openpyxl.load_workbook(CAMINHO_CONTROLE_V2, data_only=True, read_only=True)
    ws = wb[wb.sheetnames[0]]  # 1a aba = HISTORICOS (nome vem com encoding torto em read_only)

    saida = []
    vistos = set()
    for linha in ws.iter_rows(min_row=2, values_only=True):
        mkp, _motivo, raiz = linha[0], linha[1], linha[2]
        if not raiz:
            continue
        raiz = str(raiz).strip()
        if raiz in vistos:
            continue  # varias lojas (ex. CASSOL GET NET / PAGSEGURO) compartilham raiz
        vistos.add(raiz)
        banco, conta, agencia = linha[8], linha[9], linha[10]
        # _REPASSE/_COMISSAO/_FRETE/_ESTORNO/_ANUNCIO seguem sempre o
        # padrao RAIZ + sufixo (confirmado nas 93 linhas da planilha) --
        # deriva em vez de ler as colunas (uma tinha espaco sobrando).
        saida.append({
            "mkp": str(mkp).replace("\n", " ").strip() if mkp else None,
            "raiz": raiz,
            "repasse": f"{raiz}_REPASSE",
            "comissao": f"{raiz}_COMISSAO",
            "frete": f"{raiz}_FRETE",
            "estorno": f"{raiz}_ESTORNO",
            "anuncio": f"{raiz}_ANUNCIO",
            "banco": str(banco).strip() if banco else None,
            "conta": str(conta).strip() if conta else None,
            "agencia": str(agencia).strip() if agencia else None,
        })

    with open(CAMINHO_SAIDA, "w", encoding="utf-8") as f:
        json.dump(saida, f, ensure_ascii=False, indent=2)
    return len(saida)


if __name__ == "__main__":
    n = extrair()
    print(f"Salvou {n} lojas em {CAMINHO_SAIDA}")
