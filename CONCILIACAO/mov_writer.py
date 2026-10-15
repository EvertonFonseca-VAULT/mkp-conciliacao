"""Gera/acrescenta no arquivo "MOV <canal> <ddmmaaaa>.xlsx" em
C:\\GERAR MOV\\MERCADO LIVRE\\MELI_<canal>\\ -- mesmas 9 colunas que os
arquivos MOV_* que ja existem la (confirmado abrindo um real, 14/10/2026):
DATA, NATUREZA, DOCUMENTO, VENDEDOR, BANCO, AGENCIA, CONTA, HISTORICO, VALOR.

Essas 9 colunas sao um SUBCONJUNTO das linhas que ja usamos pro CMM
(rebate_cmm.py e o clipboard do dashboard, ver cmm_writer.py) -- so projeta
(ver _CAMPOS_CMM), nao recalcula nada. Arquivo NOVO (ou so acrescenta linha
no fim de um que ja existe) -- nunca edita o CMM de verdade, sem risco de
formula/validacao/Excel COM.
"""

import datetime
import os

import openpyxl

# MOV_PASTA_RAIZ_TESTE: mesma ideia de CMM_PASTA_REDE_TESTE no GUI -- aponta
# pra uma pasta de teste em vez de C:\GERAR MOV\MERCADO LIVRE de verdade.
PASTA_MOV_RAIZ = os.environ.get("MOV_PASTA_RAIZ_TESTE", r"C:\GERAR MOV\MERCADO LIVRE")

# Pasta por canal -- confirmada olhando o disco (14/10/2026). CO05 ainda nao
# tem pasta criada ali -- fica de fora de proposito (falha alto em vez de
# chutar o nome).
PASTA_POR_CANAL = {
    "FULL": "MELI_FULL", "ES": "MELI_COLETAS_ES", "PRIME": "MELI_PRIME",
    "WEB": "MELI_WEB", "MOVEIS": "MELI_MOVEIS", "ELETRO": "MELI_ELETRO",
    "CENTER": "MELI_CENTER",
}

_CABECALHO = ["DATA", "NATUREZA", "DOCUMENTO", "VENDEDOR", "BANCO", "AGENCIA", "CONTA", "HISTORICO", "VALOR"]
# Mesma ordem do cabecalho acima, so com o nome de campo usado nas linhas
# CMM (rebate_cmm.COLUNAS / dashboard clipboard, ver cmm_writer.py).
_CAMPOS_CMM = ["DATA_LACTO", "NATUREZA", "PEDIDO MKP / CV", "VENDEDOR",
               "BANCO", "AGENCIA", "CONTA", "HISTORICO", "VALOR R$"]


def _data_para_datetime(valor):
    """DATA_LACTO chega como string dd/mm/yyyy (rebate_cmm.py e o clipboard
    do dashboard so mandam texto) -- normaliza pro mesmo tipo datetime que
    os arquivos MOV existentes usam na coluna DATA."""
    if isinstance(valor, datetime.datetime):
        return valor
    if isinstance(valor, str) and valor:
        return datetime.datetime.strptime(valor, "%d/%m/%Y")
    return valor


def linhas_cmm_para_mov(linhas_cmm):
    """Projeta linhas no formato CMM (lista de dict) pros 9 campos do MOV --
    mesmo dado, outro nome de coluna. Nao faz dedupe nem calculo novo."""
    linhas_mov = []
    for linha in linhas_cmm:
        valores = [linha.get(campo) for campo in _CAMPOS_CMM]
        valores[0] = _data_para_datetime(valores[0])
        linhas_mov.append(valores)
    return linhas_mov


def caminho_mov(canal, data_ddmmaaaa):
    pasta = PASTA_POR_CANAL.get(canal)
    if not pasta:
        raise ValueError(f"Pasta MOV nao confirmada pro canal {canal} (ver mov_writer.PASTA_POR_CANAL).")
    return os.path.join(PASTA_MOV_RAIZ, pasta, f"MOV {canal} {data_ddmmaaaa}.xlsx")


def incluir_linhas_mov(caminho_xlsx, linhas_mov):
    """Cria o arquivo (com cabecalho) se nao existir; acrescenta linhas no
    fim se ja existir (pra rebate e clipboard poderem cair no mesmo arquivo
    em cliques separados do mesmo dia). Devolve a ultima linha escrita."""
    if not linhas_mov:
        return None
    if os.path.exists(caminho_xlsx):
        wb = openpyxl.load_workbook(caminho_xlsx)
        ws = wb[wb.sheetnames[0]]
        ultima = ws.max_row
    else:
        os.makedirs(os.path.dirname(caminho_xlsx), exist_ok=True)
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.append(_CABECALHO)
        ultima = 1
    for linha in linhas_mov:
        ws.append(linha)
    wb.save(caminho_xlsx)
    return ultima + len(linhas_mov)


if __name__ == "__main__":
    import shutil
    import tempfile

    with tempfile.TemporaryDirectory() as tmp:
        PASTA_MOV_RAIZ = tmp  # module-level script, sem funcao em volta -- nao precisa de global

        linhas_cmm = [{
            "DATA_LACTO": "14/10/2026", "NATUREZA": 1107, "PEDIDO MKP / CV": "123",
            "VENDEDOR": "578", "BANCO": "237", "AGENCIA": "3708", "CONTA": "X103925",
            "HISTORICO": "MKP_MELI_WE_RECEBIMENTO DE REBATE", "VALOR R$": 1.23,
        }]
        linhas_mov = linhas_cmm_para_mov(linhas_cmm)
        caminho = caminho_mov("WEB", "14102026")
        ultima = incluir_linhas_mov(caminho, linhas_mov)
        assert ultima == 2, ultima
        ultima2 = incluir_linhas_mov(caminho, linhas_mov)  # acrescenta, nao sobrescreve
        assert ultima2 == 3, ultima2

        wb = openpyxl.load_workbook(caminho)
        ws = wb[wb.sheetnames[0]]
        assert [c.value for c in ws[1]] == _CABECALHO
        linha2 = [c.value for c in ws[2]]
        assert linha2[1] == 1107 and linha2[8] == 1.23, linha2
        print("mov_writer.py: ok")
