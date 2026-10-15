from playwright.sync_api import sync_playwright
import re
import json
import openpyxl
import os
import time

CAMINHO_EXCEL_ENTRADA = r'C:\Users\everton.fonseca\Desktop\Projetos\AJUSTES.xlsx'
CAMINHO_EXCEL_SAIDA_MP = r'C:\Users\everton.fonseca\Desktop\Projetos\CONSULTAS_MP.xlsx'
CAMINHO_EXCEL_SAIDA_ML = r'C:\Users\everton.fonseca\Desktop\Projetos\CONSULTAS_ML.xlsx'


# ---------------------------------------------------------------------------
# Genérico (Excel / conexão com o Chrome)
# ---------------------------------------------------------------------------

def conectar_chrome_existente(playwright):
    browser = playwright.chromium.connect_over_cdp("http://localhost:9222")
    context = browser.contexts[0]
    return context


def limpar_numero(valor):
    if valor is None:
        return None
    if isinstance(valor, float):
        return str(int(valor))
    return str(valor).strip()


def abrir_ou_criar_planilha(caminho):
    if os.path.exists(caminho):
        wb = openpyxl.load_workbook(caminho)
        ws = wb.active
    else:
        wb = openpyxl.Workbook()
        ws = wb.active
    return wb, ws


def encontrar_bloco_pedido(ws, pedido):
    for row in range(1, ws.max_row + 1):
        if ws.cell(row=row, column=1).value == "Pedido" and str(ws.cell(row=row, column=2).value) == str(pedido):
            start = row
            end = row + 1
            while end <= ws.max_row and ws.cell(row=end, column=1).value != "Pedido":
                end += 1
            return start, end - 1
    return None


def salvar_dados(caminho, dados):
    wb, ws = abrir_ou_criar_planilha(caminho)
    linhas = dados["linhas"]
    pedido = dados["pedido"]

    # a primeira linha do bloco precisa ser sempre ("Pedido", <id>) para o
    # encontrar_bloco_pedido conseguir localizar o bloco depois
    if not linhas or linhas[0][0] != "Pedido":
        linhas = [("Pedido", pedido)] + linhas

    bloco = encontrar_bloco_pedido(ws, pedido)

    if bloco:
        start, end = bloco
        ws.delete_rows(start, end - start + 1)
        ws.insert_rows(start, len(linhas))
        for i, (label, valor) in enumerate(linhas):
            ws.cell(row=start + i, column=1).value = label
            ws.cell(row=start + i, column=2).value = valor
    else:
        if ws.max_row > 1 or ws.cell(row=1, column=1).value is not None:
            ws.append([])
        for label, valor in linhas:
            ws.append([label, valor])

    wb.save(caminho)


# ---------------------------------------------------------------------------
# Mercado Pago
# ---------------------------------------------------------------------------

def encontrar_aba_mercadopago(context):
    for page in context.pages:
        if "mercadopago.com.br" in page.url:
            return page
    return None


def extrair_lista_resultados_mp(html):
    match = re.search(r'_n\.ctx\.r=(\{.*?\});_n\.ctx\.r\.', html, re.DOTALL)
    if not match:
        return []
    data = json.loads(match.group(1))
    list_data = data["appProps"]["pageProps"].get("listData")
    if not list_data:
        return []
    itens = []
    for grupo in list_data.get("groups", []):
        for item in grupo.get("items", []):
            itens.append(item)
    return itens


def valor_amount(amount):
    if not amount:
        return None
    fraction = amount.get("fraction", "0")
    cents = amount.get("cents", "0") or "0"
    negativo = fraction.strip().startswith("-")
    fraction_limpo = fraction.replace(".", "").replace("-", "")
    inteiro = int(fraction_limpo) if fraction_limpo else 0
    decimal = int(cents) if cents else 0
    valor = inteiro + (decimal / 100)
    return round(-valor if negativo else valor, 2)


def extrair_dados_detalhe_mp(html):
    match = re.search(r'_n\.ctx\.r=(\{.*?\});_n\.ctx\.r\.', html, re.DOTALL)
    if not match:
        return None

    data = json.loads(match.group(1))
    page_props = data["appProps"]["pageProps"]
    activity = page_props["activitiesContext"]["activity"]["detail"]
    componentes = {c["id"]: c["data"] for c in page_props["detail"]["components"]}

    def texto(comp_id):
        comp = componentes.get(comp_id, {})
        return comp.get("content", {}).get("text")

    # "ticket" vem None (nao ausente) em atividades sem compra por tras (ex.
    # cashback, "resourceType": "cashback") -- o antigo .get("ticket", {})
    # nao pegava isso (a chave existe, so o valor e None) e quebrava com
    # AttributeError em .get("data", {}). Pedido do Everton, 23/09/2026.
    ticket = (componentes.get("header", {}).get("ticket") or {}).get("data") or {}

    linhas = [
        ("Pedido", activity.get("operationId")),
        ("Status", activity.get("status")),
        ("Cliente", texto("widget-client")),
        ("Produto", texto("widget-purchased-products")),
        ("Transação", texto("widget-transaction-number")),
    ]

    # Motivo (widget-reason): so aparece em atividades sem ticket de compra
    # (cashback etc, ex. "Entrada de dinheiro por seu envio"). additional_info
    # e a descricao completa, com o pedido/venda relacionado embutido (ex.
    # "Reembolsamos R$BRL2.872,45 pelo que aconteceu na devolucao de sua
    # venda 2000017769032090") -- pedido do Everton, 23/09/2026.
    reason = componentes.get("widget-reason", {})
    motivo = reason.get("content", {}).get("text")
    motivo_detalhe = reason.get("additional_info", {}).get("text")
    if motivo:
        linhas.append(("Motivo", motivo))
    if motivo_detalhe:
        linhas.append(("Motivo detalhe", motivo_detalhe))

    if ticket:
        total = ticket.get("total", {})
        linhas.append((f"{total.get('title', 'Total')} (R$)", valor_amount(total.get("amount"))))

        for grupo in ticket.get("items", {}).values():
            linhas.append((f"{grupo.get('title')} (R$)", valor_amount(grupo.get("amount"))))
            for elem in grupo.get("elements", []):
                linhas.append((f"  - {elem.get('title')} (R$)", valor_amount(elem.get("amount"))))

    return {"pedido": activity.get("operationId"), "linhas": linhas}


def buscar_e_abrir_ultimo_resultado_mp(page, numero):
    page.goto(f"https://www.mercadopago.com.br/activities?q={numero}")
    page.wait_for_load_state("networkidle")

    itens = extrair_lista_resultados_mp(page.content())
    if not itens:
        return None

    ultimo = itens[-1]
    link = ultimo.get("link")
    if not link:
        return None

    page.goto(f"https://www.mercadopago.com.br{link}")
    page.wait_for_load_state("networkidle")
    return page.content()


# ---------------------------------------------------------------------------
# Mercado Livre
# ---------------------------------------------------------------------------

def encontrar_aba_mercadolivre(context):
    for page in context.pages:
        if "mercadolivre.com.br" in page.url:
            return page
    return None


def valor_brl(texto):
    """Converte strings tipo 'R$ 595,00' ou '-R$ 62,48' em float."""
    if not texto:
        return None
    texto = texto.strip()
    negativo = texto.startswith("-")
    texto_limpo = texto.replace("-", "").replace("R$", "").strip()
    texto_limpo = texto_limpo.replace(".", "").replace(",", ".")
    if not texto_limpo:
        return None
    try:
        valor = float(texto_limpo)
    except ValueError:
        return None
    return round(-valor if negativo else valor, 2)


def extrair_pack_id_da_lista(html):
    """Pega o(s) pack id(s) do link 'ver detalhe' na lista de resultados da busca."""
    matches = re.findall(r'/vendas/(\d+)/detalhe', html)
    if not matches:
        return None
    # remove duplicados mantendo ordem e pega o último (mais recente), igual ao MP
    pack_ids = list(dict.fromkeys(matches))
    return pack_ids[-1]


def _componente_por_tipo(resp, ui_type):
    for comp in resp.values():
        if isinstance(comp, dict) and comp.get("uiType") == ui_type:
            return comp.get("data", {}) or {}
    return {}


def _texto_por_id(resp, comp_id):
    comp = resp.get(comp_id) or {}
    return (comp.get("data") or {}).get("text")


def extrair_dados_detalhe_ml(html):
    match = re.search(r'_n\.ctx\.r=(\{.*?\});_n\.ctx\.r\.', html, re.DOTALL)
    if not match:
        return None

    data = json.loads(match.group(1))
    page_props = data["appProps"]["pageProps"]
    resp = page_props.get("response") or {}
    pack_id = page_props.get("packId")
    if not pack_id:
        return None

    status = _componente_por_tipo(resp, "status")

    linhas = [
        ("Pedido", pack_id),
        ("Venda", _texto_por_id(resp, "sales_id")),
        ("Produto", _texto_por_id(resp, "detail_title")),
        ("Comprador", _texto_por_id(resp, "buyer_info_text")),
        ("Status", status.get("label")),
        ("Status detalhe", status.get("subLabel")),
    ]

    # linhas financeiras: qualquer componente "account_rows-*"
    # (PRODUCT, CHARGES, SHIPMENT, BONUSES, TOTAL, e outras que apareçam)
    for comp_id, comp in resp.items():
        if not isinstance(comp, dict) or not comp_id.startswith("account_rows-"):
            continue
        d = comp.get("data", {}) or {}
        if d.get("type") == "TOTAL":
            for row in d.get("rows", []):
                linhas.append((row.get("label", "Total") + " (R$)", valor_brl(row.get("price"))))
        else:
            titulo = d.get("title") or comp_id
            linhas.append((f"{titulo} (R$)", valor_brl(d.get("subTotal"))))
            for row in d.get("rows", []):
                linhas.append((f"  - {row.get('label')} (R$)", valor_brl(row.get("price"))))

    return {"pedido": pack_id, "linhas": linhas}


def buscar_e_abrir_ultimo_resultado_ml(page, numero):
    url = (
        "https://www.mercadolivre.com.br/vendas/omni/lista"
        f"?platform.id=ML&channel=marketshops&filters=&sort=&page=1&search={numero}"
        "&startPeriod=&toCurrent=&fromCurrent="
    )
    page.goto(url)
    page.wait_for_load_state("networkidle")

    pack_id = extrair_pack_id_da_lista(page.content())
    if not pack_id:
        return None

    page.goto(f"https://www.mercadolivre.com.br/vendas/{pack_id}/detalhe")
    page.wait_for_load_state("networkidle")
    return page.content()


# ---------------------------------------------------------------------------
# Main
# ---------------------------------------------------------------------------

def main():
    wb_entrada = openpyxl.load_workbook(CAMINHO_EXCEL_ENTRADA)
    ws_entrada = wb_entrada.active

    with sync_playwright() as p:
        context = conectar_chrome_existente(p)

        page_mp = encontrar_aba_mercadopago(context)
        page_ml = encontrar_aba_mercadolivre(context)

        for row in range(2, ws_entrada.max_row + 1):
            numero = limpar_numero(ws_entrada.cell(row=row, column=1).value)  # coluna A = PEDIDO MKP

            if numero is None:
                continue

            # números que começam com 4 são do Mercado Livre (shipping id)
            if numero.startswith("4"):
                if page_ml is None:
                    page_ml = context.new_page()
                html_detalhe = buscar_e_abrir_ultimo_resultado_ml(page_ml, numero)
                if html_detalhe is None:
                    print(f"Linha {row} | {numero} | Nada encontrado (Mercado Livre)")
                    continue

                dados = extrair_dados_detalhe_ml(html_detalhe)
                if dados:
                    salvar_dados(CAMINHO_EXCEL_SAIDA_ML, dados)
                    print(f"Linha {row} | {numero} | Salvo ML (pedido {dados['pedido']})")
                else:
                    print(f"Linha {row} | {numero} | Não consegui extrair dados do detalhe (ML)")

            else:
                if page_mp is None:
                    print(f"Linha {row} | {numero} | PULADO (nenhuma aba do Mercado Pago aberta)")
                    continue

                html_detalhe = buscar_e_abrir_ultimo_resultado_mp(page_mp, numero)
                if html_detalhe is None:
                    print(f"Linha {row} | {numero} | Nada encontrado (Mercado Pago)")
                    continue

                dados = extrair_dados_detalhe_mp(html_detalhe)
                if dados:
                    salvar_dados(CAMINHO_EXCEL_SAIDA_MP, dados)
                    print(f"Linha {row} | {numero} | Salvo MP (pedido {dados['pedido']})")
                else:
                    print(f"Linha {row} | {numero} | Não consegui extrair dados do detalhe (MP)")

            time.sleep(2)


if __name__ == "__main__":
    main()
