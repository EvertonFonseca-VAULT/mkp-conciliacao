import os
import sys

import requests
import openpyxl
from bs4 import BeautifulSoup
import tkinter as tk
from tkinter import messagebox

BASE_URL = "http://172.30.2.5:8190/sri"
CAMINHO_EXCEL = r'C:\Users\everton.fonseca\Desktop\Projetos\PROCURAR.xlsx'


def ler_credenciais():
    """Credenciais do SRI -- nunca no codigo. Prioridade: variavel de ambiente
    (mesmo padrao de KONCILI_USUARIO/KONCILI_SENHA em KONCILI/koncili.py), senao
    Gerenciador de Credenciais do Windows (keyring, servico "MARKETPLACE_SRI").
        $env:SRI_USUARIO="usuario"; $env:SRI_SENHA="senha"
    """
    usuario = os.environ.get("SRI_USUARIO")
    senha = os.environ.get("SRI_SENHA")
    if not usuario or not senha:
        import keyring
        usuario = usuario or keyring.get_password("MARKETPLACE_SRI", "usuario")
        senha = senha or keyring.get_password("MARKETPLACE_SRI", "senha")
    faltando = [n for n, v in (("SRI_USUARIO", usuario), ("SRI_SENHA", senha)) if not v]
    if faltando:
        sys.exit(f"ERRO: defina a(s) variavel(is) de ambiente ou grave no cofre (MARKETPLACE_SRI): {', '.join(faltando)}")
    return usuario, senha


def login(session, usuario, senha):
    payload = f'LoginForm%5Busername%5D={usuario}&LoginForm%5Bpassword%5D={senha}&LoginForm%5BrememberMe%5D=0&yt0=Login'
    headers = {'Content-Type': 'application/x-www-form-urlencoded'}
    resp = session.post(f"{BASE_URL}/site/login", headers=headers, data=payload)
    resp.raise_for_status()
    return resp


def inicializar_pesquisa(session):
    resp = session.get(f"{BASE_URL}/comercial/pedidos/pesquisa")
    resp.raise_for_status()
    return resp


def buscar_pedido(session, pedmkp):
    payload = {
        'txtfilial': '', 'txtpedido': '', 'txtdatade': '', 'txtdataate': '',
        'txtvendedor': '', 'cgc': '', 'txtpedsite': '', 'pedmkp': pedmkp,
        'serie': '', 'nota': '', 'yt0': 'Pesquisa'
    }
    headers = {'Content-Type': 'application/x-www-form-urlencoded'}
    resp = session.post(f"{BASE_URL}/comercial/pedidos/pesquisa", headers=headers, data=payload)
    resp.raise_for_status()
    resp.encoding = 'cp1252'
    return resp


def extrair_dados_pedido(html):
    soup = BeautifulSoup(html, "html.parser")
    linha = soup.find("tr", onclick=True)
    if linha is None:
        return None
    tds = linha.find_all("td")
    valores = [td.get_text(strip=True) for td in tds]
    campos = ["filial", "pedido", "emissao", "cliente", "nro_doc",
              "nota", "frete", "vendedor", "ped_site", "ped_mkp"]
    return dict(zip(campos, valores))


def limpar_pedmkp(valor):
    if valor is None:
        return None
    if isinstance(valor, float):
        return str(int(valor))
    return str(valor).strip()


def buscar_e_validar(session, pedmkp_buscado):
    pedmkp_limpo = limpar_pedmkp(pedmkp_buscado)

    inicializar_pesquisa(session)  # repete antes de cada busca (necessário pra não trazer resultado "fantasma")

    resp = buscar_pedido(session, pedmkp_limpo)
    dados = extrair_dados_pedido(resp.text)

    if dados is None:
        return None, resp

    if dados["ped_mkp"] != pedmkp_limpo:
        return None, resp

    return dados, resp


def processar_planilha(caminho, usuario, senha):
    wb = openpyxl.load_workbook(caminho)
    ws = wb.active

    session = requests.Session()
    login(session, usuario, senha)
    inicializar_pesquisa(session)

    total = 0
    encontrados = 0

    for row in range(2, ws.max_row + 1):
        valor_a = ws.cell(row=row, column=1).value
        valor_b = ws.cell(row=row, column=2).value

        if valor_b is None:
            continue

        total += 1
        dados, _ = buscar_e_validar(session, valor_b)

        if dados is not None:
            ws.cell(row=row, column=3).value = valor_b
            encontrados += 1
        else:
            ws.cell(row=row, column=3).value = valor_a

    wb.save(caminho)
    return total, encontrados


if __name__ == "__main__":
    usuario, senha = ler_credenciais()
    total, encontrados = processar_planilha(CAMINHO_EXCEL, usuario, senha)

    root = tk.Tk()
    root.withdraw()
    messagebox.showinfo(
        "Processamento concluído",
        f"Total processado: {total}\nEncontrados: {encontrados}\nNão encontrados: {total - encontrados}"
    )